"""
生成需求工程流水线 Skills 总结 PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 颜色常量 ──
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x2E)
DARK_BLUE = RGBColor(0x16, 0x21, 0x3E)
ACCENT_BLUE = RGBColor(0x0F, 0x4C, 0x81)
ACCENT_TEAL = RGBColor(0x00, 0x96, 0x88)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
LIGHT_BLUE = RGBColor(0xE3, 0xF2, 0xFD)
GRAY = RGBColor(0x75, 0x75, 0x75)
MEDIUM_GRAY = RGBColor(0xBD, 0xBD, 0xBD)
TABLE_HEADER_BG = RGBColor(0x0F, 0x4C, 0x81)
TABLE_ROW_ALT = RGBColor(0xF0, 0xF4, 0xF8)
ORANGE = RGBColor(0xFF, 0x98, 0x00)
GREEN = RGBColor(0x4C, 0xAF, 0x50)
RED = RGBColor(0xE5, 0x39, 0x35)

FONT_CN = '微软雅黑'
FONT_EN = 'Segoe UI'
FONT_MONO = 'Consolas'


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=14,
                 bold=False, color=BLACK, alignment=PP_ALIGN.LEFT,
                 font_name=FONT_CN, line_spacing=1.2):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    p.line_spacing = Pt(font_size * line_spacing)
    return txBox, tf


def add_paragraph(tf, text, font_size=14, bold=False, color=BLACK,
                  alignment=PP_ALIGN.LEFT, font_name=FONT_CN,
                  space_before=0, space_after=0, line_spacing=1.3):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = Pt(font_size * line_spacing)
    return p


def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def style_table(table, header_bg=TABLE_HEADER_BG, alt_row=TABLE_ROW_ALT):
    """统一表格样式"""
    for col_idx in range(len(table.columns)):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_bg
        for p in cell.text_frame.paragraphs:
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.name = FONT_CN
            p.alignment = PP_ALIGN.CENTER

    for row_idx in range(1, len(table.rows)):
        for col_idx in range(len(table.columns)):
            cell = table.cell(row_idx, col_idx)
            if row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = alt_row
            else:
                cell.fill.background()
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.name = FONT_CN
                p.font.color.rgb = BLACK
                p.alignment = PP_ALIGN.LEFT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_section_title(slide, text, top=0.3):
    add_text_box(slide, 0.6, top, 8.5, 0.5, text,
                 font_size=22, bold=True, color=ACCENT_BLUE)
    # 下划线装饰
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.6), Inches(top + 0.55), Inches(1.2), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT_TEAL
    line.line.fill.background()


# ══════════════════════════════════════════
# 开始构建 PPT
# ══════════════════════════════════════════
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ─────────────────────────────────────────
# Slide 1: 封面
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide, DARK_BLUE)

# 顶部装饰条
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                              Inches(0), Inches(0), Inches(13.333), Inches(0.08))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT_TEAL
bar.line.fill.background()

# 主标题
add_text_box(slide, 1.5, 1.8, 10, 1.2,
             '需求工程流水线 Skills 总结',
             font_size=40, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

# 副标题
add_text_box(slide, 1.5, 3.2, 10, 0.8,
             'HarmonyOS ArkTS 增量开发 — 从原始需求到标准化 Spec',
             font_size=20, color=ACCENT_TEAL, alignment=PP_ALIGN.CENTER)

# 分隔线
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                               Inches(4.5), Inches(4.3), Inches(4.333), Inches(0.03))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_TEAL
line.line.fill.background()

# 底部信息
add_text_box(slide, 1.5, 4.8, 10, 0.5,
             'AntennaPod ArkTS Project  |  KAIZEN-Migration Methodology',
             font_size=14, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1.5, 5.3, 10, 0.5,
             '2026-04-13',
             font_size=13, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)

# ─────────────────────────────────────────
# Slide 2: 目录
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, '目录')

items = [
    ('01', 'Pipeline 整体架构', '四步流水线的设计思路与数据流'),
    ('02', '四个 Skill 概览', '每个 Skill 的职责、输入输出、核心产出'),
    ('03', 'rq-parse — 需求解析', '意图分类、领域映射、模糊点检测'),
    ('04', 'rq-clarify — 需求澄清', '交互式提问、HarmonyOS 专项澄清'),
    ('05', 'rq2spec — 需求转 Spec', 'Spec 标准格式、技术约束自动注入'),
    ('06', 'spec-check — Spec 校验', '五维校验模型、评分体系'),
    ('07', '关键设计决策', '兼容性、门控、自动注入等核心决策'),
    ('08', '触发示例 & 下游衔接', '如何触发、与 spec-evolver 衔接'),
]

for i, (num, title, desc) in enumerate(items):
    y = 1.2 + i * 0.7
    # 编号圆点
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                     Inches(1.0), Inches(y + 0.05), Inches(0.4), Inches(0.4))
    circle.fill.solid()
    circle.fill.fore_color.rgb = ACCENT_BLUE
    circle.line.fill.background()
    tf = circle.text_frame
    tf.paragraphs[0].text = num
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.name = FONT_EN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    circle.text_frame.word_wrap = False

    add_text_box(slide, 1.7, y, 4, 0.3, title,
                 font_size=15, bold=True, color=BLACK)
    add_text_box(slide, 1.7, y + 0.3, 7, 0.3, desc,
                 font_size=11, color=GRAY)

# ─────────────────────────────────────────
# Slide 3: Pipeline 架构
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, 'Pipeline 整体架构')

# 流程图 - 4个阶段方块 + 箭头
stages = [
    ('rq-parse', '需求解析', '意图分类\n领域映射\n模糊点标记', ACCENT_BLUE),
    ('rq-clarify', '需求澄清', '交互式提问\n消除歧义\n补齐信息', RGBColor(0x00, 0x79, 0x6B)),
    ('rq2spec', '需求转 Spec', 'Spec 模板\n约束注入\n门控确认', RGBColor(0xE6, 0x51, 0x00)),
    ('spec-check', 'Spec 校验', '五维校验\nPASS/WARN/FAIL\n自动修复', RGBColor(0x6A, 0x1B, 0x9A)),
]

box_w = 2.3
box_h = 1.8
gap = 0.7
start_x = 1.2
y_top = 2.0

for i, (name, title, desc, color) in enumerate(stages):
    x = start_x + i * (box_w + gap)

    # 主方块
    shape = add_rounded_rect(slide, x, y_top, box_w, box_h, color)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = name
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.name = FONT_MONO
    tf.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xCC)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT_CN
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(6)

    for line_text in desc.split('\n'):
        p = tf.add_paragraph()
        p.text = line_text
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
        p.font.name = FONT_CN
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(2)

    # 箭头 (除最后一个)
    if i < len(stages) - 1:
        arrow_x = x + box_w
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            Inches(arrow_x + 0.1), Inches(y_top + box_h / 2 - 0.15),
            Inches(gap - 0.2), Inches(0.3)
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = MEDIUM_GRAY
        arrow.line.fill.background()

# 输入输出标注
add_text_box(slide, 0.2, y_top + box_h / 2 - 0.2, 1.0, 0.4,
             '用户\n原始需求', font_size=10, color=GRAY, alignment=PP_ALIGN.CENTER)

end_x = start_x + 3 * (box_w + gap) + box_w + 0.2
add_text_box(slide, end_x, y_top + 0.1, 1.5, 0.6,
             '校验通过', font_size=11, bold=True, color=GREEN, alignment=PP_ALIGN.LEFT)

# 下游箭头
down_shape = add_rounded_rect(slide, end_x, y_top + 0.9, 2.0, 0.7, RGBColor(0xE8, 0xF5, 0xE9))
tf = down_shape.text_frame
tf.paragraphs[0].text = 'arkts-spec-evolver'
tf.paragraphs[0].font.size = Pt(10)
tf.paragraphs[0].font.name = FONT_MONO
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = GREEN
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
p = tf.add_paragraph()
p.text = '进入实现阶段'
p.font.size = Pt(9)
p.font.color.rgb = GRAY
p.font.name = FONT_CN
p.alignment = PP_ALIGN.CENTER

# 底部说明
add_text_box(slide, 1.2, 4.5, 9, 1.0,
             '数据流：每个 Skill 的产出是下一个 Skill 的输入。需求清晰时可跳过 rq-clarify 直接进入 rq2spec。\n'
             'Spec 格式兼容现有 arkts-spec-evolver 增量 Spec 体系，生成后可直接纳入 spec/ 目录管理。',
             font_size=11, color=GRAY, line_spacing=1.5)

# ─────────────────────────────────────────
# Slide 4: 四个 Skill 概览表
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, '四个 Skill 概览')

# 表格
rows_data = [
    ['Skill', '职责', '核心产出', '参考文档', '触发词示例'],
    ['rq-parse', '解析原始需求：意图分类、领域映射、\n影响层次、关键实体、模糊点标记',
     '需求解析报告\n（对话输出）', 'domain-taxonomy.md\n领域分类体系', '"我想加个功能"\n"能不能支持XX"'],
    ['rq-clarify', '交互式澄清：分批提问、选项驱动、\nHarmonyOS 专项问题、提前终止',
     '澄清后需求摘要\n（含决策表）', 'question-bank.md\n按领域分类问题库', '"帮我完善需求"\n"还缺什么"'],
    ['rq2spec', '需求转 Spec：8 大章节模板、\n技术约束自动注入、HARD-GATE 确认',
     '标准化 Spec 文件\n（.md）', 'spec-template.md\n完整模板+写作指南', '"生成 spec"\n"文档化需求"'],
    ['spec-check', '五维校验：结构/内容/可行性/\n一致性/对齐，自动修复建议',
     '校验报告\n（PASS/WARN/FAIL）', 'validation-rules.md\n30+ 条校验规则', '"检查 spec"\n"校验一下"'],
]

tbl = slide.shapes.add_table(len(rows_data), 5,
                              Inches(0.5), Inches(1.2),
                              Inches(12.3), Inches(5.0)).table
col_widths = [1.5, 3.5, 2.2, 2.5, 2.6]
for i, w in enumerate(col_widths):
    tbl.columns[i].width = Inches(w)

for r, row in enumerate(rows_data):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

# 首列加粗
for r in range(1, len(rows_data)):
    cell = tbl.cell(r, 0)
    for p in cell.text_frame.paragraphs:
        p.font.bold = True
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_BLUE

style_table(tbl)

# ─────────────────────────────────────────
# Slide 5: rq-parse 详解
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, 'rq-parse — 需求解析')

# 左侧: 6步流程
add_text_box(slide, 0.6, 1.2, 5, 0.4, '六步工作流程',
             font_size=16, bold=True, color=ACCENT_BLUE)

steps = [
    ('Step 1', '意图分类', 'feature / enhancement / bugfix / optimization / refactoring'),
    ('Step 2', '领域映射', 'UI、数据、网络、媒体、导航、状态管理、系统能力等 9 大领域'),
    ('Step 3', '影响层次', 'View → ViewModel → Model → 数据层 → 基础设施（MVVM）'),
    ('Step 4', '实体提取', '数据实体、页面实体、接口实体、事件实体'),
    ('Step 5', '模糊点检测', 'blocker / important / nice-to-have 三级严重度'),
    ('Step 6', '复杂度评估', 'S（简单）/ M（中等）/ L（复杂）/ XL（特大）'),
]

for i, (step, title, desc) in enumerate(steps):
    y = 1.8 + i * 0.8
    # step 标签
    tag = add_rounded_rect(slide, 0.6, y, 0.8, 0.35, ACCENT_BLUE)
    tf = tag.text_frame
    tf.paragraphs[0].text = step
    tf.paragraphs[0].font.size = Pt(9)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.name = FONT_EN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text_box(slide, 1.6, y - 0.02, 2, 0.3, title,
                 font_size=13, bold=True, color=BLACK)
    add_text_box(slide, 1.6, y + 0.28, 4.5, 0.4, desc,
                 font_size=10, color=GRAY)

# 右侧: 模糊点分类表
add_text_box(slide, 7.2, 1.2, 5, 0.4, '模糊点分类',
             font_size=16, bold=True, color=ACCENT_BLUE)

fuzzy_data = [
    ['类别', '示例', '影响'],
    ['范围模糊', '"更好的发现体验"', '无法确定功能边界'],
    ['行为模糊', '"支持离线"', '无法确定实现方案'],
    ['交互模糊', '"添加搜索"', '无法设计 UI'],
    ['数据模糊', '"保存播放记录"', '无法设计数据模型'],
    ['约束模糊', '"要快"', '无法定义验收标准'],
    ['兼容模糊', '"不影响现有功能"', '无法确定回归范围'],
]

tbl = slide.shapes.add_table(len(fuzzy_data), 3,
                              Inches(7.2), Inches(1.8),
                              Inches(5.5), Inches(3.5)).table
tbl.columns[0].width = Inches(1.5)
tbl.columns[1].width = Inches(2.2)
tbl.columns[2].width = Inches(1.8)

for r, row in enumerate(fuzzy_data):
    for c, val in enumerate(row):
        tbl.cell(r, c).text = val
        tbl.cell(r, c).vertical_anchor = MSO_ANCHOR.MIDDLE
style_table(tbl)

# ─────────────────────────────────────────
# Slide 6: rq-clarify 详解
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, 'rq-clarify — 需求澄清')

# 左侧: 核心策略
add_text_box(slide, 0.6, 1.2, 5, 0.4, '核心策略',
             font_size=16, bold=True, color=ACCENT_BLUE)

strategies = [
    ('提问优先级', 'blocker → important → nice-to-have\n按严重程度排序，优先解决阻塞性问题'),
    ('批次控制', '每轮最多 3-5 个问题\n相关问题合并，依赖问题串联'),
    ('选项驱动', '尽量提供可选方案让用户选择\n降低回答难度，加快决策速度'),
    ('提前终止', '所有 blocker 解决后可终止\n未澄清的点标记为 assumption'),
]

for i, (title, desc) in enumerate(strategies):
    y = 1.8 + i * 1.15
    shape = add_rounded_rect(slide, 0.6, y, 5.5, 0.95, LIGHT_BLUE)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = title
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ACCENT_BLUE
    tf.paragraphs[0].font.name = FONT_CN

    for line in desc.split('\n'):
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(10)
        p.font.color.rgb = BLACK
        p.font.name = FONT_CN
        p.space_before = Pt(2)

# 右侧: 提问模板
add_text_box(slide, 7.0, 1.2, 5, 0.4, '六种提问模板',
             font_size=16, bold=True, color=ACCENT_BLUE)

templates = [
    ('范围模糊', '边界界定法', '列出可能方面让用户选择'),
    ('行为模糊', '场景还原法', '给出具体场景确认预期行为'),
    ('交互模糊', '原型描述法', '提供两种实现方案对比'),
    ('数据模糊', '字段枚举法', '列出字段清单确认增减'),
    ('约束模糊', '量化锚定法', '提供基线和目标让用户定标'),
    ('兼容模糊', '影响枚举法', '列出受影响功能确认范围'),
]

tbl = slide.shapes.add_table(len(templates) + 1, 3,
                              Inches(7.0), Inches(1.8),
                              Inches(5.8), Inches(3.5)).table
tbl.columns[0].width = Inches(1.5)
tbl.columns[1].width = Inches(1.8)
tbl.columns[2].width = Inches(2.5)

headers = ['模糊点类别', '提问方法', '核心思路']
for c, h in enumerate(headers):
    tbl.cell(0, c).text = h
for r, (cat, method, idea) in enumerate(templates):
    tbl.cell(r + 1, 0).text = cat
    tbl.cell(r + 1, 1).text = method
    tbl.cell(r + 1, 2).text = idea
for r in range(len(templates) + 1):
    for c in range(3):
        tbl.cell(r, c).vertical_anchor = MSO_ANCHOR.MIDDLE
style_table(tbl)

# HarmonyOS 专项
add_text_box(slide, 7.0, 5.6, 5.8, 0.4, 'HarmonyOS 专项澄清',
             font_size=13, bold=True, color=RGBColor(0x00, 0x79, 0x6B))
add_text_box(slide, 7.0, 6.0, 5.8, 0.8,
             '数据存储方案选择（Preferences vs RelationalStore vs fileIo）\n'
             '离线/在线策略  |  后台行为  |  权限需求与降级方案',
             font_size=10, color=GRAY)

# ─────────────────────────────────────────
# Slide 7: rq2spec 详解
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, 'rq2spec — 需求转标准化 Spec')

# 左侧: Spec 模板结构
add_text_box(slide, 0.6, 1.2, 5, 0.4, 'Spec 标准格式（8 大章节）',
             font_size=16, bold=True, color=ACCENT_BLUE)

sections = [
    ('YAML', 'Frontmatter', 'id / type / title / priority / status / affects'),
    ('1', '背景与动机', '当前状态、驱动力、价值主张'),
    ('2', '目标', '主要目标（1-3个）+ 非目标（范围排除）'),
    ('3', '功能需求', 'FR-N: 描述 + 用户故事 + 交互流程 + 边界条件'),
    ('4', '非功能需求', 'NFR-N: 性能 / 兼容性 / 可用性 / 数据完整性'),
    ('5', '验收标准', 'AC-N: Given-When-Then 格式，可测试'),
    ('6', '约束', '技术约束 / 业务约束 / 设计约束'),
    ('7', '影响分析', '涉及模块 + 依赖关系 + 风险评估'),
    ('8', '附录', '待定假设 + 参考资料 + 术语表'),
]

for i, (num, title, desc) in enumerate(sections):
    y = 1.8 + i * 0.55
    tag_color = RGBColor(0xE6, 0x51, 0x00) if num != 'YAML' else ACCENT_TEAL
    tag_w = 0.6 if num != 'YAML' else 0.7
    tag = add_rounded_rect(slide, 0.6, y, tag_w, 0.3, tag_color)
    tf = tag.text_frame
    tf.paragraphs[0].text = num
    tf.paragraphs[0].font.size = Pt(9)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.name = FONT_EN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text_box(slide, 1.5, y - 0.02, 2, 0.25, title,
                 font_size=11, bold=True, color=BLACK)
    add_text_box(slide, 3.5, y - 0.02, 3.5, 0.25, desc,
                 font_size=9, color=GRAY)

# 右侧: 技术约束自动注入
add_text_box(slide, 7.2, 1.2, 5, 0.4, '技术约束自动注入',
             font_size=16, bold=True, color=ACCENT_BLUE)

inject_data = [
    ['涉及领域', '自动注入约束'],
    ['音视频播放', 'AVPlayer 状态机顺序、fd:// 协议、后台三要素'],
    ['文件下载', 'request.agent API（>1MB）、沙箱路径存储'],
    ['数据库', 'INSERT OR REPLACE 改 id、ResultSet close'],
    ['导航', 'NavDestination 参数 onReady 获取'],
    ['网络请求', 'http.createHttp() 必须 destroy、INTERNET 权限'],
    ['系统图标', '只使用已验证的 sys.symbol 名称'],
]

tbl = slide.shapes.add_table(len(inject_data), 2,
                              Inches(7.2), Inches(1.8),
                              Inches(5.5), Inches(3.2)).table
tbl.columns[0].width = Inches(1.8)
tbl.columns[1].width = Inches(3.7)
for r, row in enumerate(inject_data):
    for c, val in enumerate(row):
        tbl.cell(r, c).text = val
        tbl.cell(r, c).vertical_anchor = MSO_ANCHOR.MIDDLE
style_table(tbl)

# 门控说明
add_text_box(slide, 7.2, 5.3, 5.5, 0.4, 'HARD-GATE 门控',
             font_size=14, bold=True, color=RED)
gate_box = add_rounded_rect(slide, 7.2, 5.8, 5.5, 1.2, RGBColor(0xFF, 0xEB, 0xEE))
tf = gate_box.text_frame
tf.word_wrap = True
tf.paragraphs[0].text = 'Spec 生成后必须经用户确认才能进入校验：'
tf.paragraphs[0].font.size = Pt(10)
tf.paragraphs[0].font.color.rgb = BLACK
tf.paragraphs[0].font.name = FONT_CN
for action in ['确认 → 触发 spec-check 校验',
               '要求修改 → 修改后重新确认',
               '跳过校验 → 直接存入 spec/ 目录']:
    p = tf.add_paragraph()
    p.text = '  ' + action
    p.font.size = Pt(9)
    p.font.color.rgb = GRAY
    p.font.name = FONT_CN
    p.space_before = Pt(2)

# ─────────────────────────────────────────
# Slide 8: spec-check 详解
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, 'spec-check — Spec 校验')

# 五维模型
add_text_box(slide, 0.6, 1.2, 12, 0.4, '五维校验模型',
             font_size=16, bold=True, color=ACCENT_BLUE)

dims = [
    ('D1', '结构完整性', 'Structure', 'frontmatter 字段\n必需章节\n编号连续', ACCENT_BLUE),
    ('D2', '内容质量', 'Content', '背景非空话\nAC 覆盖 FR\nGWT 格式', RGBColor(0x00, 0x79, 0x6B)),
    ('D3', '技术可行性', 'Feasibility', 'API 存在性\n已知踩坑回避\n权限声明', RGBColor(0xE6, 0x51, 0x00)),
    ('D4', '内部一致性', 'Consistency', '目标↔FR 对齐\n非目标不冲突\naffects 匹配', RGBColor(0x6A, 0x1B, 0x9A)),
    ('D5', '项目对齐', 'Alignment', 'id 唯一性\n目录结构\nArkTS 严格模式', RGBColor(0x42, 0x42, 0x42)),
]

for i, (code, title, eng, checks, color) in enumerate(dims):
    x = 0.6 + i * 2.5
    shape = add_rounded_rect(slide, x, 1.8, 2.2, 2.6, color)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = code
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xCC)
    tf.paragraphs[0].font.name = FONT_EN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT_CN
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = eng
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    p.font.name = FONT_EN
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(2)

    for line in checks.split('\n'):
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        p.font.name = FONT_CN
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(4)

# 评分体系
add_text_box(slide, 0.6, 4.8, 12, 0.4, '评分体系',
             font_size=16, bold=True, color=ACCENT_BLUE)

scores = [
    ('PASS', '通过', '所有必需检查项通过', GREEN),
    ('WARN', '警告', '可选检查项未通过，不阻塞', ORANGE),
    ('FAIL', '不通过', '必需检查项未通过，必须修复', RED),
]

for i, (label, cn, desc, color) in enumerate(scores):
    x = 0.6 + i * 4.2
    tag = add_rounded_rect(slide, x, 5.3, 1.0, 0.4, color)
    tf = tag.text_frame
    tf.paragraphs[0].text = label
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.name = FONT_EN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text_box(slide, x + 1.1, 5.3, 3, 0.2, cn,
                 font_size=12, bold=True, color=color)
    add_text_box(slide, x + 1.1, 5.55, 3, 0.2, desc,
                 font_size=10, color=GRAY)

add_text_box(slide, 0.6, 6.1, 12, 0.5,
             '整体判定：全部 PASS/WARN → Spec 可进入实现阶段；任一 FAIL → 必须修复后重新校验',
             font_size=11, bold=True, color=BLACK)

# ─────────────────────────────────────────
# Slide 9: 关键设计决策
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, '关键设计决策')

decisions = [
    ('Spec 格式兼容',
     'frontmatter 的 id / type / status / affects 结构完全兼容 arkts-spec-evolver，'
     '生成后可无缝衔接进入实现阶段，无需人工格式转换。'),
    ('HARD-GATE 门控',
     'rq2spec 生成 Spec 后设有门控，需用户确认后才触发 spec-check。'
     '防止错误的 Spec 直接进入下游，同时给用户修改机会。'),
    ('模糊点三级严重度',
     'blocker（不澄清无法写 Spec）/ important（影响质量）/ nice-to-have（可带入实现阶段），'
     '避免过度追问导致用户疲惫。'),
    ('技术约束自动注入',
     'rq2spec 根据涉及的技术领域，自动在约束章节注入 CLAUDE.md 中 Top-10 已知踩坑，'
     '防止开发者重复踩坑。'),
    ('五维独立评分',
     'spec-check 五个维度独立评分（PASS/WARN/FAIL），'
     '区分"必须修复"和"建议改进"，不一刀切。'),
]

for i, (title, desc) in enumerate(decisions):
    y = 1.2 + i * 1.15
    # 编号
    num_shape = add_rounded_rect(slide, 0.6, y + 0.05, 0.45, 0.45, ACCENT_BLUE)
    tf = num_shape.text_frame
    tf.paragraphs[0].text = str(i + 1)
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.name = FONT_EN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text_box(slide, 1.3, y, 11, 0.35, title,
                 font_size=15, bold=True, color=BLACK)
    add_text_box(slide, 1.3, y + 0.4, 11, 0.6, desc,
                 font_size=11, color=GRAY, line_spacing=1.4)

# ─────────────────────────────────────────
# Slide 10: 触发示例 & 下游衔接
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, WHITE)
add_section_title(slide, '触发示例 & 下游衔接')

# 触发表
add_text_box(slide, 0.6, 1.2, 5, 0.4, '触发示例',
             font_size=16, bold=True, color=ACCENT_BLUE)

trigger_data = [
    ['用户说', '触发 Skill'],
    ['"我想给播客添加睡眠定时器"', 'rq-parse'],
    ['"能不能支持离线下载"', 'rq-parse'],
    ['"这个需求还缺什么信息"', 'rq-clarify'],
    ['"帮我完善一下这个需求"', 'rq-clarify'],
    ['"生成 spec"', 'rq2spec'],
    ['"把需求转成标准化格式"', 'rq2spec'],
    ['"检查一下这个 spec"', 'spec-check'],
    ['"spec 写得对不对"', 'spec-check'],
]

tbl = slide.shapes.add_table(len(trigger_data), 2,
                              Inches(0.6), Inches(1.8),
                              Inches(5.5), Inches(4.5)).table
tbl.columns[0].width = Inches(3.5)
tbl.columns[1].width = Inches(2.0)
for r, row in enumerate(trigger_data):
    for c, val in enumerate(row):
        tbl.cell(r, c).text = val
        tbl.cell(r, c).vertical_anchor = MSO_ANCHOR.MIDDLE
        if r > 0 and c == 1:
            for p in tbl.cell(r, c).text_frame.paragraphs:
                p.font.name = FONT_MONO
                p.font.bold = True
                p.font.color.rgb = ACCENT_BLUE
style_table(tbl)

# 下游衔接
add_text_box(slide, 7.0, 1.2, 5.5, 0.4, '与现有体系衔接',
             font_size=16, bold=True, color=ACCENT_BLUE)

connections = [
    ('上游', '用户原始需求（自然语言）', '本 pipeline 的起点'),
    ('下游', 'arkts-spec-evolver', 'Spec 确认后进入执行管理'),
    ('验证', 'arkts-knowledge-verifier', '校验 ArkTS API 可行性'),
    ('定位', 'arkts-codebase-locator', '判断影响范围和现有代码'),
    ('执行', 'superpowers:writing-plans', '生成实现计划'),
]

for i, (role, target, desc) in enumerate(connections):
    y = 1.8 + i * 1.0
    tag = add_rounded_rect(slide, 7.0, y, 1.0, 0.35, ACCENT_TEAL)
    tf = tag.text_frame
    tf.paragraphs[0].text = role
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.name = FONT_CN
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text_box(slide, 8.2, y - 0.02, 4, 0.25, target,
                 font_size=12, bold=True, color=BLACK, font_name=FONT_MONO)
    add_text_box(slide, 8.2, y + 0.25, 4, 0.25, desc,
                 font_size=10, color=GRAY)

# ─────────────────────────────────────────
# Slide 11: 文件结构总览
# ─────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BLUE)

# 装饰条
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                              Inches(0), Inches(0), Inches(13.333), Inches(0.08))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT_TEAL
bar.line.fill.background()

add_text_box(slide, 0.6, 0.5, 12, 0.5, '文件结构总览',
             font_size=24, bold=True, color=WHITE)

# 四个 skill 的文件结构
skill_files = [
    ('rq-parse/', [
        'SKILL.md — 需求解析主文档',
        'evals.json — 3 个测试用例',
        'references/domain-taxonomy.md — 领域分类体系'
    ]),
    ('rq-clarify/', [
        'SKILL.md — 需求澄清主文档',
        'evals.json — 3 个测试用例',
        'references/question-bank.md — 按领域分类问题库'
    ]),
    ('rq2spec/', [
        'SKILL.md — 需求转 Spec 主文档',
        'evals.json — 3 个测试用例',
        'references/spec-template.md — 完整模板+写作指南'
    ]),
    ('spec-check/', [
        'SKILL.md — Spec 校验主文档',
        'evals.json — 3 个测试用例',
        'references/validation-rules.md — 30+ 条校验规则'
    ]),
]

for i, (folder, files) in enumerate(skill_files):
    x = 0.6 + i * 3.1
    shape = add_rounded_rect(slide, x, 1.3, 2.8, 3.5, RGBColor(0x1E, 0x2D, 0x4A))
    tf = shape.text_frame
    tf.word_wrap = True

    tf.paragraphs[0].text = folder
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ACCENT_TEAL
    tf.paragraphs[0].font.name = FONT_MONO
    tf.paragraphs[0].space_after = Pt(8)

    for f in files:
        p = tf.add_paragraph()
        p.text = f
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        p.font.name = FONT_CN
        p.space_before = Pt(6)

# 底部统计
add_text_box(slide, 0.6, 5.3, 12, 0.5,
             '总计：4 个 Skill  |  12 个文件  |  ~950 行 SKILL.md  |  ~550 行 references  |  12 个测试用例',
             font_size=14, bold=True, color=ACCENT_TEAL, alignment=PP_ALIGN.CENTER)

# Thank you
add_text_box(slide, 1.5, 6.0, 10, 0.8,
             'Ready to use — 在对话中直接输入需求即可触发流水线',
             font_size=16, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)

# ── 保存 ──
output_path = os.path.join(os.path.dirname(__file__),
                           'RequirementEngineering_Skills_Summary.pptx')
prs.save(output_path)
print(f'PPT saved to: {output_path}')
